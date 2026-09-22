import os
import zipfile
import xml.etree.ElementTree as ET
import pandas as pd
from openpyxl import load_workbook
from pathlib import Path

def parse_kml_elementos(kml_content):
    """Extrai OLT, PON e os componentes do KMZ, ignorando itens 'EXISTENTE'."""
    tree = ET.ElementTree(ET.fromstring(kml_content))
    root = tree.getroot()
    namespace = {'kml': 'http://www.opengis.net/kml/2.2'}
    
    dados_processados = []
    
    def varrer_pastas(elemento, olt_atual=None, pon_atual=None):
        for folder in elemento.findall('kml:Folder', namespace):
            nome_folder = folder.find('kml:name', namespace)
            txt_folder = nome_folder.text if nome_folder is not None else ""
            
            if "OLT" in txt_folder.upper() and "PON" in txt_folder.upper():
                partes = txt_folder.split()
                for p in partes:
                    if p.upper().startswith("OLT"):
                        olt_atual = p.replace("OLT-", "").replace("OLT", "").strip()
                    elif "PON" in p.upper():
                        pon_atual = p.replace("PON", "").strip()
            
            for placemark in folder.findall('kml:Placemark', namespace):
                name_elem = placemark.find('kml:name', namespace)
                if name_elem is not None and name_elem.text:
                    nome_item = name_elem.text.strip()
                    if "EXISTENTE" in nome_item.upper():
                        continue
                    dados_processados.append({
                        'OLT': olt_atual,
                        'PON': pon_atual,
                        'Item': nome_item
                    })
            varrer_pastas(folder, olt_atual, pon_atual)

    varrer_pastas(root)
    return dados_processados

def categorizar_itens(lista_itens):
    cxs, ces, sps, cxes = [], [], [], []
    ce_hub = ""
    linhas_necessarias = 1
    
    for item in lista_itens:
        txt = item['Item']
        upper_txt = txt.upper()
        
        if 'SP-' in upper_txt:
            sps.append(txt)
            if '1/8' in upper_txt:
                linhas_necessarias = max(linhas_necessarias, 8)
            elif '1/4' in upper_txt:
                linhas_necessarias = max(linhas_necessarias, 4)
            elif '1/16' in upper_txt:
                linhas_necessarias = max(linhas_necessarias, 16)
        elif upper_txt.startswith('CE-') or 'CE-' in upper_txt:
            ces.append(txt)
        elif 'CXE' in upper_txt:
            cxes.append(txt)
        elif upper_txt.startswith('CX-') or 'CX-' in upper_txt:
            cxs.append(txt)
            
    if ces:
        ce_hub = ces[0]
        
    linhas_necessarias = max(linhas_necessarias, len(cxs))

    return {
        'olt': lista_itens[0]['OLT'] if lista_itens else '',
        'pon': lista_itens[0]['PON'] if lista_itens else '',
        'CX': cxs,
        'CE': ces,
        'SP': sps,
        'CE_HUB': ce_hub,
        'linhas_necessarias': linhas_necessarias
    }

def mapear_colunas_excel(ws):
    mapeamento = {}
    linha_cabecalho = 3
    
    for r in range(1, 6):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(row=r, column=c).value).strip().upper()
            if 'CX' in val and ('NVO' in val or 'NOVO' in val):
                mapeamento['CX'] = c
            elif 'CE' in val and ('NVO' in val or 'NOVO' in val) and 'HUB' not in val:
                mapeamento['CE'] = c
            elif 'SP' in val and 'HUB' in val:
                mapeamento['SP'] = c
            elif 'CE' in val and 'HUB' in val:
                mapeamento['CE_HUB'] = c
            elif val == 'OLT':
                mapeamento['OLT'] = c
            elif 'PON' in val:
                mapeamento['PON'] = c
                
        if len(mapeamento) >= 4:
            linha_cabecalho = r
            break
            
    return mapeamento, linha_cabecalho

def encontrar_arquivo_kmz():
    nome_input = input("Digite o nome do arquivo KMZ (ex: AHS - JOSE LUIS BERNARDES.kmz): ").strip()
    if not nome_input.lower().endswith('.kmz'):
        nome_input += '.kmz'
        
    if os.path.exists(nome_input):
        return nome_input
        
    caminho_downloads = Path.home() / "Downloads" / nome_input
    if caminho_downloads.exists():
        return str(caminho_downloads)
        
    print(f"Aviso: O arquivo '{nome_input}' não foi encontrado na pasta atual nem em Downloads.")
    return None

def processar_kmz_para_planilha(caminho_excel):
    caminho_kmz = encontrar_arquivo_kmz()
    if not caminho_kmz:
        return

    elementos_brutos = []
    with zipfile.ZipFile(caminho_kmz, 'r') as kmz:
        for arquivo in kmz.namelist():
            if arquivo.endswith('.kml'):
                with kmz.open(arquivo) as kml_file:
                    elementos_brutos.extend(parse_kml_elementos(kml_file.read()))
                    
    if not elementos_brutos:
        print("Nenhum elemento válido encontrado no KMZ (ou continham 'EXISTENTE').")
        return

    dados = categorizar_itens(elementos_brutos)
    print(f"Lido do KMZ -> OLT: {dados['olt']} | PON: {dados['pon']} | Linhas necessárias: {dados['linhas_necessarias']}")
    
    wb = load_workbook(caminho_excel)
    
    # Procura em TODAS as abas da planilha automaticamente
    aba_encontrada = None
    linhas_alvo = []
    cols = {}
    linha_cabecalho = 3
    
    for nome_aba in wb.sheetnames:
        ws = wb[nome_aba]
        c, l_cab = mapear_colunas_excel(ws)
        
        if 'OLT' in c and 'PON' in c:
            # Verifica se essa OLT e PON existem nesta aba
            l_alvo_temp = []
            for row in range(l_cab + 1, ws.max_row + 1):
                val_olt = str(ws.cell(row=row, column=c['OLT']).value).strip()
                val_pon = str(ws.cell(row=row, column=c['PON']).value).strip()
                
                if val_olt == str(dados['olt']) and val_pon == str(dados['pon']):
                    l_alvo_temp.append(row)
            
            if l_alvo_temp:
                aba_encontrada = nome_aba
                linhas_alvo = l_alvo_temp
                cols = c
                linha_cabecalho = l_cab
                break
                
    if not aba_encontrada:
        print(f"Aviso: O bloco OLT {dados['olt']} / PON {dados['pon']} não foi encontrado em nenhuma aba da planilha.")
        return
        
    print(f"Bloco encontrado na aba: '{aba_encontrada}'!")
    ws = wb[aba_encontrada]

    # Preenche os dados dinamicamente
    for idx, row in enumerate(linhas_alvo):
        if idx >= dados['linhas_necessarias']:
            break
            
        if idx < len(dados['CX']) and 'CX' in cols:
            ws.cell(row=row, column=cols['CX'], value=dados['CX'][idx])
            
        if idx < len(dados['CE']) and 'CE' in cols:
            ws.cell(row=row, column=cols['CE'], value=dados['CE'][idx])
            
        if 'SP' in cols and dados['SP']:
            sp_texto = " / ".join(dados['SP'])
            ws.cell(row=row, column=cols['SP'], value=sp_texto)
            
        if 'CE_HUB' in cols and dados['CE_HUB']:
            ws.cell(row=row, column=cols['CE_HUB'], value=dados['CE_HUB'])
            
    wb.save(caminho_excel)
    print(f"Sucesso! Dados preenchidos na aba '{aba_encontrada}' para OLT {dados['olt']} PON {dados['pon']}.")

# --- EXECUÇÃO ---
if __name__ == "__main__":
    planilha_excel = "Planilha Certificação de Rede GPON.xlsx"
    processar_kmz_para_planilha(planilha_excel)